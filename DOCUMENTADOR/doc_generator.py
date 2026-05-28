"""
Generacion y exportacion de documentos desde datos de proyecto.
Formatos: Markdown, DOCX, HTML, PDF.
"""
from __future__ import annotations

import re
import textwrap
from pathlib import Path
from datetime import datetime


# ─── Filtrado de secciones por modo ─────────────────────────────────────────

def _get_filtered_groups(project: dict, mode: str) -> dict:
    """Devuelve {grupo: [seccion_id]} filtrado segun modo: Markdown/Tecnico/No Tecnico."""
    from project_manager import get_sections_for_type
    ptype   = project.get("type", "Web")
    profile = project.get("profile", "Ambos")
    all_groups = get_sections_for_type(ptype, profile)
    if mode == "Tecnico":
        return {k: v for k, v in all_groups.items() if "Vision Funcional" not in k}
    if mode == "No Tecnico":
        return {k: v for k, v in all_groups.items() if "Vision Funcional" in k}
    return all_groups


# ─── Generador Markdown ──────────────────────────────────────────────────────

def _md_content(project: dict, mode: str = "Markdown") -> str:
    lines: list[str] = []
    name    = project.get("name", "Proyecto")
    ptype   = project.get("type", "")
    stack   = project.get("stack", "")
    desc    = project.get("desc", "")
    profile = project.get("profile", "")
    date    = datetime.now().strftime("%d/%m/%Y")

    lines += [f"# {name}", "",
              f"**Tipo:** {ptype}  |  **Perfil:** {profile}  |  **Fecha:** {date}"]
    if stack:
        lines.append(f"**Stack:** {stack}")
    lines += ["", "---", ""]

    if desc:
        lines += ["## Descripcion", "", desc, "", "---", ""]

    sections_data: dict = project.get("sections", {})

    for group, section_ids in _get_filtered_groups(project, mode).items():
        lines += [f"## {group}", ""]
        for sid in section_ids:
            content = sections_data.get(sid, "").strip()
            lines += [f"### {sid}", ""]
            lines.append(content if content else "_[Pendiente de completar]_")
            lines.append("")

    return "\n".join(lines)


# ─── API pública ─────────────────────────────────────────────────────────────

def render_preview(project: dict, mode: str = "Markdown") -> str:
    return _md_content(project, mode)


def generate_all(project: dict) -> None:
    """Guarda MD + HTML en carpeta exports/ junto al exe (o junto al .py en dev)."""
    import sys
    name    = project.get("name", "proyecto")
    base    = Path(sys.executable).parent if getattr(sys, "frozen", False) \
              else Path(__file__).resolve().parent
    out_dir = base / "exports"
    out_dir.mkdir(exist_ok=True)
    _export_md(  project, str(out_dir / f"{name}.md"))
    _export_html(project, str(out_dir / f"{name}.html"))


def export_document(project: dict, path: str, fmt: str) -> None:
    if "Markdown" in fmt:
        _export_md(project, path)
    elif "Word" in fmt:
        _export_docx(project, path)
    elif "HTML" in fmt:
        _export_html(project, path)
    elif "PDF" in fmt:
        _export_pdf(project, path)
    else:
        _export_md(project, path)


# ─── Exportador MD ───────────────────────────────────────────────────────────

def _export_md(project: dict, path: str) -> None:
    Path(path).write_text(_md_content(project), encoding="utf-8")


# ─── Exportador HTML ─────────────────────────────────────────────────────────

def _export_html(project: dict, path: str) -> None:
    try:
        import markdown as md_lib
        exts = ["tables", "fenced_code"]
        try:
            md_lib.markdown("", extensions=["nl2br"])
            exts.append("nl2br")
        except Exception:
            pass
        html_body = md_lib.markdown(_md_content(project), extensions=exts)
    except ImportError:
        html_body = f"<pre>{_md_content(project)}</pre>"

    name    = project.get("name", "Documento")
    ptype   = project.get("type", "")
    profile = project.get("profile", "")
    date    = datetime.now().strftime("%d/%m/%Y")

    html = textwrap.dedent(f"""\
        <!DOCTYPE html>
        <html lang="es">
        <head>
          <meta charset="UTF-8">
          <meta name="viewport" content="width=device-width, initial-scale=1">
          <title>{name} — Documentacion</title>
          <style>
            * {{ box-sizing: border-box; }}
            body {{ font-family: 'Segoe UI', Arial, sans-serif;
                    max-width: 980px; margin: 0 auto;
                    padding: 32px 20px; color: #212121;
                    line-height: 1.65; background: #F0F2F5; }}
            .cover {{ background: #0D47A1; color: white;
                       padding: 32px 40px; border-radius: 6px;
                       margin-bottom: 32px; }}
            .cover h1 {{ margin: 0 0 8px; font-size: 2rem;
                          border: none; color: white; }}
            .cover .meta {{ font-size: 0.9rem; opacity: 0.85; }}
            .content {{ background: white; padding: 32px 40px;
                        border-radius: 6px;
                        box-shadow: 0 1px 4px rgba(0,0,0,.08); }}
            h1 {{ color: #0D47A1; border-bottom: 2px solid #1565C0;
                  padding-bottom: 8px; margin-top: 40px; }}
            h2 {{ color: #1565C0; margin-top: 36px;
                  border-left: 4px solid #1565C0;
                  padding-left: 12px; }}
            h3 {{ color: #37474F; margin-top: 24px; }}
            hr {{ border: none; border-top: 1px solid #CFD8DC; margin: 28px 0; }}
            pre {{ background: #F0F2F5; padding: 16px 20px;
                   border-radius: 4px; overflow-x: auto; font-size: .88rem; }}
            code {{ background: #ECEFF1; padding: 2px 6px;
                    border-radius: 3px; font-size: .88rem; }}
            pre code {{ background: none; padding: 0; }}
            table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
            th {{ background: #E3F2FD; color: #1565C0; font-weight: 600; }}
            th, td {{ border: 1px solid #CFD8DC; padding: 8px 14px; text-align: left; }}
            tr:nth-child(even) td {{ background: #F8F9FA; }}
            blockquote {{ border-left: 4px solid #90CAF9; margin: 12px 0;
                          padding: 8px 16px; background: #E3F2FD;
                          color: #37474F; border-radius: 0 4px 4px 0; }}
            ul, ol {{ padding-left: 24px; }}
            li {{ margin: 4px 0; }}
            .footer {{ text-align: center; font-size: .8rem;
                       color: #90A4AE; margin-top: 40px; }}
          </style>
        </head>
        <body>
          <div class="cover">
            <h1>{name}</h1>
            <div class="meta">Tipo: {ptype}  ·  Perfil: {profile}  ·  Generado: {date}</div>
          </div>
          <div class="content">
        {html_body}
          </div>
          <p class="footer">Generado con Documentador de Proyectos</p>
        </body>
        </html>
    """)
    Path(path).write_text(html, encoding="utf-8")


# ─── Exportador DOCX ─────────────────────────────────────────────────────────

def _export_docx(project: dict, path: str) -> None:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as e:
        raise RuntimeError("python-docx no instalado. Ejecuta: pip install python-docx") from e

    doc = Document()

    # Margenes
    for sec in doc.sections:
        sec.top_margin    = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin   = Inches(1.2)
        sec.right_margin  = Inches(1.2)

    def _shading(paragraph, fill_hex: str) -> None:
        try:
            pPr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"),   "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"),  fill_hex)
            pPr.append(shd)
        except Exception:
            pass

    def _add_inline(para, line: str) -> None:
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", line):
            if part.startswith("**") and part.endswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*") and part.endswith("*"):
                run = para.add_run(part[1:-1])
                run.italic = True
            elif part.startswith("`") and part.endswith("`"):
                run = para.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif part:
                para.add_run(part)

    def _flush_code(buf: list[str]) -> None:
        if not buf:
            return
        cp = doc.add_paragraph()
        run = cp.add_run("\n".join(buf))
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        _shading(cp, "F0F2F5")

    def _process_content(content: str) -> None:
        in_code = False
        code_buf: list[str] = []
        for line in content.split("\n"):
            if line.startswith("```"):
                if in_code:
                    _flush_code(code_buf)
                    code_buf = []
                    in_code  = False
                else:
                    in_code = True
                continue
            if in_code:
                code_buf.append(line)
                continue
            if line.startswith("### "):
                h = doc.add_heading(line[4:], level=3)
                if h.runs: h.runs[0].font.color.rgb = RGBColor(0x37, 0x47, 0x4F)
            elif line.startswith("## "):
                h = doc.add_heading(line[3:], level=3)
                if h.runs: h.runs[0].font.color.rgb = RGBColor(0x15, 0x65, 0xC0)
            elif line.startswith("# "):
                h = doc.add_heading(line[2:], level=3)
                if h.runs: h.runs[0].font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
            elif line.startswith("> "):
                bp = doc.add_paragraph(line[2:])
                bp.paragraph_format.left_indent = Inches(0.4)
            elif line.startswith(("- ", "* ")):
                try:
                    doc.add_paragraph(line[2:], style="List Bullet")
                except Exception:
                    doc.add_paragraph(f"• {line[2:]}")
            elif line.strip() in ("", "---"):
                doc.add_paragraph("")
            else:
                _add_inline(doc.add_paragraph(), line)
        if in_code:
            _flush_code(code_buf)

    # ── Portada
    tp = doc.add_heading(project.get("name", "Proyecto"), level=0)
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if tp.runs:
        tp.runs[0].font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        f"Tipo: {project.get('type', '')}  ·  "
        f"Perfil: {project.get('profile', '')}  ·  "
        f"Fecha: {datetime.now().strftime('%d/%m/%Y')}"
    )
    mr.font.color.rgb = RGBColor(0x60, 0x7D, 0x8B)
    mr.font.size = Pt(10)

    if project.get("stack"):
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sr = sp.add_run(f"Stack: {project.get('stack', '')}")
        sr.font.color.rgb = RGBColor(0x60, 0x7D, 0x8B)
        sr.font.size = Pt(10)

    doc.add_page_break()

    # ── Descripcion
    desc = project.get("desc", "")
    if desc:
        dh = doc.add_heading("Descripcion General", level=1)
        if dh.runs:
            dh.runs[0].font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)
        doc.add_paragraph(desc)

    # ── Grupos y secciones
    sections_data: dict = project.get("sections", {})

    for group, section_ids in _get_filtered_groups(project, "Markdown").items():
        gh = doc.add_heading(group, level=1)
        if gh.runs:
            gh.runs[0].font.color.rgb = RGBColor(0x0D, 0x47, 0xA1)

        for sid in section_ids:
            sh = doc.add_heading(sid, level=2)
            if sh.runs:
                sh.runs[0].font.color.rgb = RGBColor(0x15, 0x65, 0xC0)

            content = sections_data.get(sid, "").strip()
            if content:
                _process_content(content)
            else:
                ip = doc.add_paragraph("[Pendiente de completar]")
                if ip.runs:
                    ip.runs[0].font.color.rgb = RGBColor(0x90, 0xA4, 0xAE)
                    ip.runs[0].italic = True

    doc.save(path)


# ─── Exportador PDF ───────────────────────────────────────────────────────────

def _export_pdf(project: dict, path: str) -> None:
    try:
        from fpdf import FPDF
    except ImportError as e:
        raise RuntimeError("fpdf2 no instalado. Ejecuta: pip install fpdf2") from e

    def _clean(text: str) -> str:
        cleaned: list[str] = []
        in_code = False
        for line in text.split("\n"):
            if line.startswith("```"):
                in_code = not in_code
                if in_code:
                    cleaned.append("[bloque de codigo]")
                continue
            if in_code:
                continue
            line = re.sub(r"^#{1,3}\s+", "", line)
            line = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
            line = re.sub(r"\*([^*]+)\*",   r"\1", line)
            line = re.sub(r"`([^`]+)`",      r"\1", line)
            line = re.sub(r"^---$", "─" * 30, line)
            cleaned.append(line)
        return "\n".join(cleaned)

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    # Portada
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(13, 71, 161)
    pdf.multi_cell(0, 12, project.get("name", "Proyecto"), align="C")
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(96, 125, 143)
    pdf.multi_cell(
        0, 8,
        f"Tipo: {project.get('type', '')}  |  Perfil: {project.get('profile', '')}  |  "
        f"Fecha: {datetime.now().strftime('%d/%m/%Y')}",
    )
    if project.get("stack"):
        pdf.multi_cell(0, 8, f"Stack: {project.get('stack', '')}")

    pdf.set_draw_color(207, 216, 220)
    pdf.line(10, pdf.get_y() + 2, 200, pdf.get_y() + 2)
    pdf.ln(8)

    desc = project.get("desc", "")
    if desc:
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(21, 101, 192)
        pdf.multi_cell(0, 10, "Descripcion")
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(33, 33, 33)
        pdf.multi_cell(0, 7, desc)
        pdf.ln(4)

    sections_data: dict = project.get("sections", {})

    for group, section_ids in _get_filtered_groups(project, "Markdown").items():
        pdf.set_font("Helvetica", "B", 13)
        pdf.set_text_color(13, 71, 161)
        pdf.multi_cell(0, 11, group)
        pdf.set_draw_color(207, 216, 220)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(3)

        for sid in section_ids:
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(21, 101, 192)
            pdf.multi_cell(0, 9, sid)

            content = sections_data.get(sid, "").strip()
            if content:
                pdf.set_font("Helvetica", "", 9)
                pdf.set_text_color(33, 33, 33)
                pdf.multi_cell(0, 6, _clean(content))
            else:
                pdf.set_font("Helvetica", "I", 9)
                pdf.set_text_color(144, 164, 174)
                pdf.multi_cell(0, 7, "[Pendiente de completar]")
            pdf.ln(3)

    pdf.output(path)
